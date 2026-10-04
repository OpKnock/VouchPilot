"""Unit tests for messy reader (fixtures built with openpyxl, no files)."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from vouch_engine import normalise
from vouch_engine.messy import read_messy_xlsx


def _save_merged_title(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.merge_cells("A1:F1")
    ws["A1"] = "SHARMA TRADERS - SALES REGISTER"
    # Row 2 left blank.
    headers = ["Seller", "Buyer", "Invoice No", "Date", "Amount"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=h)
    ws.merge_cells("A4:A5")
    ws.cell(row=4, column=1, value="Sharma Traders")
    ws.cell(row=4, column=2, value="Nagpur Agro")
    ws.cell(row=4, column=3, value="SI/001")
    ws.cell(row=4, column=4, value="31/03/2026")
    ws.cell(row=4, column=5, value=50000)
    # A5 is merged; only set the other columns.
    ws.cell(row=5, column=2, value="Pune Foods")
    ws.cell(row=5, column=3, value="SI/002")
    ws.cell(row=5, column=4, value="01/04/2026")
    ws.cell(row=5, column=5, value=25000)
    wb.save(str(path))


def test_merged_title_and_merged_seller(tmp_path: Path) -> None:
    xlsx = tmp_path / "merged.xlsx"
    _save_merged_title(xlsx)
    result = read_messy_xlsx(str(xlsx))
    headers = result["headers"]
    assert "Seller" in headers
    assert len(result["rows"]) == 2
    assert result["rows"][0]["Seller"] == "Sharma Traders"
    assert result["rows"][1]["Seller"] == "Sharma Traders"
    blob = " ".join(str(v) for r in result["rows"] for v in r.values())
    assert "SALES REGISTER" not in blob
    assert result["report"]["merged_ranges_filled"] >= 2
    assert result["report"]["header_row"] == 3


def _save_multi_sheet(path: Path) -> None:
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.cell(row=1, column=1, value="junk")
    ws_sum.cell(row=1, column=2, value="junk2")
    ws_sum.cell(row=2, column=1, value="x")
    ws_sum.cell(row=2, column=2, value="y")
    ws_bills = wb.create_sheet("Bills")
    headers = ["Seller", "Buyer", "Invoice No", "Date", "Amount"]
    for col, h in enumerate(headers, start=1):
        ws_bills.cell(row=1, column=col, value=h)
    data = [
        ["Sharma", "Nagpur", "SI/001", "31/03/2026", 100],
        ["Sharma", "Pune", "SI/002", "01/04/2026", 200],
        ["Sharma", "Delhi", "SI/003", "02/04/2026", 300],
    ]
    for r, row in enumerate(data, start=2):
        for c, v in enumerate(row, start=1):
            ws_bills.cell(row=r, column=c, value=v)
    wb.save(str(path))


def test_multi_sheet_picks_bills(tmp_path: Path) -> None:
    xlsx = tmp_path / "multi.xlsx"
    _save_multi_sheet(xlsx)
    best = read_messy_xlsx(str(xlsx))
    assert best["report"]["sheet_used"] == "Bills"
    assert len(best["rows"]) == 3
    names = {t["name"] for t in best["report"]["sheets_tried"]}
    assert {"Summary", "Bills"} <= names

    junk = read_messy_xlsx(str(xlsx), sheet="Summary")
    assert junk["report"]["sheet_used"] == "Summary"
    assert "junk" in junk["headers"]

    with pytest.raises(KeyError) as exc:
        read_messy_xlsx(str(xlsx), sheet="Nope")
    msg = str(exc.value)
    assert "Summary" in msg and "Bills" in msg


def test_hindi_headers_map_downstream(tmp_path: Path) -> None:
    xlsx = tmp_path / "hindi.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    hindi = ["विक्रेता", "खरीदार", "बीजक", "तारीख", "राशि"]
    for col, h in enumerate(hindi, start=1):
        ws.cell(row=1, column=col, value=h)
    ws.cell(row=2, column=1, value="Sharma Traders")
    ws.cell(row=2, column=2, value="Nagpur Agro")
    ws.cell(row=2, column=3, value="SI/001")
    ws.cell(row=2, column=4, value="31/03/2026")
    ws.cell(row=2, column=5, value=50000)
    ws.cell(row=3, column=1, value="Sharma Traders")
    ws.cell(row=3, column=2, value="Pune Foods")
    ws.cell(row=3, column=3, value="SI/002")
    ws.cell(row=3, column=4, value="01/04/2026")
    ws.cell(row=3, column=5, value=25000)
    wb.save(str(xlsx))

    result = read_messy_xlsx(str(xlsx))
    headers = result["headers"]
    mapping = normalise.map_columns(headers, result["rows"])
    # Translated headers must hit the canonical fields.
    assert mapping[headers[0]] == "seller.name"
    assert mapping[headers[1]] == "buyer.name"
    assert mapping[headers[2]] == "doc.invoice_number"
    assert mapping[headers[3]] == "doc.date"
    assert mapping[headers[4]] == "money.total"
    assert "seller.name" in mapping.values()


def test_determinism_and_empty_sheet(tmp_path: Path) -> None:
    xlsx = tmp_path / "det.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="Seller")
    ws.cell(row=1, column=2, value="Amount")
    ws.cell(row=2, column=1, value="A")
    ws.cell(row=2, column=2, value=10)
    wb.save(str(xlsx))

    first = read_messy_xlsx(str(xlsx))
    second = read_messy_xlsx(str(xlsx))
    assert first == second

    empty_path = tmp_path / "empty.xlsx"
    wb2 = Workbook()
    wb2.active.title = "Sheet1"
    wb2.save(str(empty_path))
    empty = read_messy_xlsx(str(empty_path))
    assert empty["headers"] == []
    assert empty["rows"] == []

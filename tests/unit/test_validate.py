"""Unit tests for prediction validation and writers."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from vouch_engine.schemas import Prediction
from vouch_engine.validate import validate_prediction, write_jsonl, write_xlsx


def _good(i: int = 1) -> dict:
    return {
        "row_id": i,
        "invoice_number": f"SI/{i:03d}",
        "voucher_type": "Sales",
        "confidence": 0.97,
        "needs_review": False,
        "top_k": [["Sales", 0.97]],
        "evidence": ["PERSPECTIVE: seller"],
    }


def test_validate_prediction_ok() -> None:
    pred = validate_prediction(_good())
    assert isinstance(pred, Prediction)
    assert pred.voucher_type == "Sales"


def test_validate_prediction_rejects_bad_label() -> None:
    with pytest.raises(Exception):
        validate_prediction({**_good(), "voucher_type": "Not A Label"})


def test_write_jsonl_counts_invalid(tmp_path: Path) -> None:
    out = tmp_path / "preds.jsonl"
    stats = write_jsonl([_good(1), {**_good(2), "voucher_type": "Bogus"}, _good(3)], out)
    assert stats == {"written": 2, "invalid": 1}
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == 2
    assert json.loads(lines[0])["voucher_type"] == "Sales"


def test_write_xlsx_counts_invalid(tmp_path: Path) -> None:
    out = tmp_path / "preds.xlsx"
    stats = write_xlsx([_good(1), {"row_id": 0, "invoice_number": "", "voucher_type": "Sales"}], out)
    assert stats == {"written": 1, "invalid": 1}
    wb = load_workbook(str(out))
    ws = wb.active
    assert ws is not None
    assert ws.max_row == 2  # header + 1 valid row
    assert ws.cell(row=2, column=3).value == "Sales"
    wb.close()


def test_writers_never_crash_on_garbage(tmp_path: Path) -> None:
    assert write_jsonl(["garbage", None, 42], tmp_path / "a.jsonl")["written"] == 0
    assert write_xlsx(["garbage", None, 42], tmp_path / "b.xlsx")["written"] == 0

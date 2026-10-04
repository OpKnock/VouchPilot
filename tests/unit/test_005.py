"""Feature 005 wiring: intake CLI, messy pipeline, API document upload."""

from fastapi.testclient import TestClient
from openpyxl import Workbook

from vouch_engine.__main__ import main
from vouch_engine import ingest, normalise
from vouch_engine.api import create_app


def _hindi_sheet(path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Sharma Traders Sales Register"
    ws.append([])
    ws.append(["विक्रेता", "खरीदार", "बीजक", "तारीख", "विवरण", "मात्रा", "राशि"])
    ws.append(["Sharma Traders", "Nagpur Agro", "SI/1", "2026-08-01", "chairs", 10, 50000])
    wb.save(path)


def test_hindi_twin_maps_like_english(tmp_path):
    from vouch_engine import messy

    hp = str(tmp_path / "hi.xlsx")
    _hindi_sheet(hp)
    got = messy.read_messy_xlsx(hp)
    assert got["rows"], "hindi rows dropped"
    mapping = normalise.map_columns(got["headers"], got["rows"])
    canon = {v for v in mapping.values() if v}
    for key in ("seller.name", "buyer.name", "doc.invoice_number", "doc.date", "money.total"):
        assert key in canon, (key, mapping)


def test_messy_title_and_merged_flow(tmp_path):
    from vouch_engine import messy

    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:D1")
    ws["A1"] = "Monthly Report"
    ws.append(["Seller", "Buyer", "Invoice No", "Taxable Value"])
    ws.append(["Sharma Traders", "X", "SI/1", 1000])
    ws.append(["", "Y", "SI/2", 2000])
    ws.merge_cells("A3:A4")
    path = str(tmp_path / "m.xlsx")
    wb.save(path)
    got = messy.read_messy_xlsx(path)
    invoiced = [r for r in got["rows"] if str(r.get("Invoice No", "")).startswith("SI/")]
    assert len(invoiced) == 2
    assert all(r.get("Seller") == "Sharma Traders" for r in invoiced)
    assert all("Monthly Report" not in str(v) for r in got["rows"] for v in r.values())


def test_intake_cli_text_pdf(tmp_path):
    import fitz

    pdf = str(tmp_path / "inv.pdf")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "TAX INVOICE SI/26-27/0412\nSeller: Sharma Traders\nTotal Rs 59000\nCGST 4500")
    doc.save(pdf)
    doc.close()
    out = str(tmp_path / "rows.xlsx")
    assert main(["intake", "--input", pdf, "--out", out]) == 0
    data = ingest.read_excel(out)
    assert len(data["rows"]) >= 1
    blob = " ".join(str(v) for r in data["rows"] for v in r.values())
    assert "SI/26-27/0412" in blob and "59000" in blob


def test_api_predict_pdf_upload(tmp_path):
    import fitz

    pdf = tmp_path / "inv.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "TAX INVOICE SI/9\nTotal Rs 1000")
    doc.save(str(pdf))
    doc.close()
    client = TestClient(create_app())
    resp = client.post("/predict", files={"file": ("inv.pdf", pdf.read_bytes())},
                       params={"scorer": "keyword"})
    assert resp.status_code == 200
    assert resp.json()["n_rows"] >= 1


def test_api_rejects_garbage_file():
    client = TestClient(create_app())
    resp = client.post("/predict", files={"file": ("x.pdf", b"not a pdf at all")},
                       params={"scorer": "keyword"})
    assert resp.status_code in (200, 400)
    if resp.status_code == 200:
        assert resp.json()["n_rows"] >= 1

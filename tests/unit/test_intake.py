"""Unit tests for vouch_engine.intake (offline, tmp_path only).

OCR-dependent assertions run only when the tesseract binary is present;
otherwise the loud-error contract (IntakeError mentioning 'winget') is asserted.
"""

from __future__ import annotations

from pathlib import Path

import fitz  # PyMuPDF
import pytest
from openpyxl import load_workbook
from PIL import Image

from vouch_engine.intake import (
    IntakeError,
    detect_type,
    intake_to_xlsx,
    parse_invoice_fields,
    pdf_page_to_image,
    pdf_to_text,
    tesseract_available,
)

SAMPLE_TEXT = """Tax Invoice
Invoice No: INV-2024-001
Date: 12/08/2024
Seller: Sharma Traders, Nagpur
Buyer: Gupta Stores, Pune
GSTIN: 27ABCDE1234F1Z5
Taxable Amount Rs 10,000
CGST Rs 900
SGST Rs 900
Total Amount: Rs 11,800
Dhanyavaad धन्यवाद"""

# ASCII twin for embedding in PDFs (builtins fonts have no Devanagari glyphs).
SAMPLE_ASCII = SAMPLE_TEXT.replace(" धन्यवाद", "")


def _make_text_pdf(path: Path, text: str = SAMPLE_ASCII) -> Path:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    doc.save(str(path))
    doc.close()
    return path


def _make_blank_pdf(path: Path) -> Path:
    doc = fitz.open()
    doc.new_page()
    doc.save(str(path))
    doc.close()
    return path


def test_detect_type_mapping() -> None:
    assert detect_type("a.xlsx") == "xlsx"
    assert detect_type("a.XLSM") == "xlsx"
    assert detect_type("a.xls") == "xlsx"
    assert detect_type("a.csv") == "csv"
    assert detect_type("a.PDF") == "pdf"
    assert detect_type("a.png") == "image"
    assert detect_type("a.JPG") == "image"
    assert detect_type("a.jpeg") == "image"
    assert detect_type("a.tiff") == "image"
    assert detect_type("a.bmp") == "image"
    assert detect_type("a.webp") == "image"


def test_detect_type_unknown_raises() -> None:
    with pytest.raises(IntakeError):
        detect_type("a.docx")
    with pytest.raises(IntakeError):
        detect_type("noextension")


def test_parse_invoice_fields_sample() -> None:
    row = parse_invoice_fields(SAMPLE_TEXT)
    assert row["invoice_number"] == "INV-2024-001"
    assert row["gstin"] == "27ABCDE1234F1Z5"
    assert row["date"] == "12/08/2024"
    assert row["total"] == "11,800"
    assert row["taxable"] == "10,000"
    assert row["cgst"] == "900"
    assert row["sgst"] == "900"
    assert "Sharma Traders" in row["seller"]
    assert "Gupta Stores" in row["buyer"]
    assert row["currency"] == "INR"
    assert "narration" in row


def test_parse_invoice_fields_missing_keys_absent() -> None:
    row = parse_invoice_fields("hello world")
    assert "invoice_number" not in row
    assert "gstin" not in row
    assert "total" not in row
    assert row["currency"] == "INR"
    assert row["narration"] == "hello world"


def test_pdf_to_text_blank_is_empty(tmp_path: Path) -> None:
    pdf = _make_blank_pdf(tmp_path / "blank.pdf")
    assert pdf_to_text(str(pdf)) == ""


def test_pdf_page_to_image_returns_pil_image(tmp_path: Path) -> None:
    pdf = _make_text_pdf(tmp_path / "inv.pdf")
    img = pdf_page_to_image(str(pdf), page=0, dpi=200)
    assert isinstance(img, Image.Image)
    assert img.width > 0 and img.height > 0


def test_text_pdf_roundtrip_to_xlsx(tmp_path: Path) -> None:
    pdf = _make_text_pdf(tmp_path / "inv.pdf")
    out = tmp_path / "out.xlsx"
    report = intake_to_xlsx(str(pdf), str(out))
    assert report["kind"] == "pdf"
    assert report["pages"] == 1
    assert report["ocr_used"] is False
    assert report["rows"] >= 1
    assert out.exists()
    wb = load_workbook(str(out))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "invoice_number" in headers
    assert ws.max_row >= 2  # header + at least one data row
    wb.close()


def test_scanned_pdf_ocr_or_loud_error(tmp_path: Path) -> None:
    pdf = _make_blank_pdf(tmp_path / "scanned.pdf")
    out = tmp_path / "out.xlsx"
    if tesseract_available():
        report = intake_to_xlsx(str(pdf), str(out))
        assert report["ocr_used"] is True
        assert report["rows"] >= 1
        assert out.exists()
    else:
        with pytest.raises(IntakeError, match="winget"):
            intake_to_xlsx(str(pdf), str(out))


def test_legacy_xls_intake_is_rejected_with_actionable_error(tmp_path):
    from vouch_engine.intake import IntakeError, intake_to_xlsx

    source = tmp_path / "legacy.xls"
    source.write_bytes(b"not-a-real-xls")

    try:
        intake_to_xlsx(str(source), str(tmp_path / "out.xlsx"))
    except IntakeError as exc:
        assert "legacy .xls" in str(exc).lower()
        assert ".xlsx" in str(exc).lower()
        return
    raise AssertionError("legacy .xls should require conversion to .xlsx")

"""CLI integration test (Team-C, US2): gold -> run(stub) -> evaluate.

The stub scorer is not expected to match gold labels, so this asserts report
shape (all keys, n_rows), metric reproducibility (evaluate twice identical),
and gold seed reproducibility (labels json + parsed xlsx rows identical).
"""

import json
import os
import subprocess
import sys

from openpyxl import Workbook, load_workbook

from vouch_engine.evaluate import REPORT_KEYS
from vouch_engine.gold import build_dataset

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(ROOT, "src")


def _run_cli(*cli_args):
    env = dict(os.environ)
    env["PYTHONPATH"] = SRC + os.pathsep + env.get("PYTHONPATH", "")
    return subprocess.run(
        [sys.executable, "-m", "vouch_engine", *cli_args],
        cwd=ROOT,
        env=env,
        capture_output=True,
        text=True,
        timeout=300,
    )


def _read_sheet_rows(xlsx_path):
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = workbook.active
    return [[c.value for c in row] for row in sheet.iter_rows()]


def test_cli_gold_run_evaluate_reproducible(tmp_path):
    prefix = str(tmp_path / "gold")
    build_dataset(60, 1, prefix)
    xlsx = prefix + ".xlsx"
    labels_json = prefix + "_labels.json"
    pred = str(tmp_path / "predictions.jsonl")

    run = _run_cli("run", "--input", xlsx, "--out", pred, "--scorer", "stub")
    assert run.returncode == 0, run.stderr
    assert "n_rows=60" in run.stdout
    with open(pred, encoding="utf-8") as handle:
        lines = [line for line in handle if line.strip()]
    assert len(lines) == 60

    rep1 = str(tmp_path / "eval1.json")
    rep2 = str(tmp_path / "eval2.json")
    eval1 = _run_cli("evaluate", "--gold", labels_json, "--pred", pred,
                     "--report", rep1)
    assert eval1.returncode == 0, eval1.stderr
    eval2 = _run_cli("evaluate", "--gold", labels_json, "--pred", pred,
                     "--report", rep2)
    assert eval2.returncode == 0, eval2.stderr

    report1 = json.loads(open(rep1, encoding="utf-8").read())
    report2 = json.loads(open(rep2, encoding="utf-8").read())
    assert set(report1) == set(REPORT_KEYS)
    assert report1["n_rows"] == 60
    assert report1 == report2, "metrics must reproduce exactly"


def test_gold_seed_reproducibility_files(tmp_path):
    first = str(tmp_path / "a" / "gold")
    second = str(tmp_path / "b" / "gold")
    build_dataset(60, 1, first)
    build_dataset(60, 1, second)
    labels1 = json.loads(open(first + "_labels.json", encoding="utf-8").read())
    labels2 = json.loads(open(second + "_labels.json", encoding="utf-8").read())
    assert labels1 == labels2
    # xlsx bytes may differ by embedded timestamp: compare parsed rows instead.
    assert _read_sheet_rows(first + ".xlsx") == _read_sheet_rows(second + ".xlsx")


def test_cli_run_uses_resilient_reader_for_messy_workbook(tmp_path):
    xlsx = str(tmp_path / "messy.xlsx")
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "VouchPilot export"
    tx = wb.create_sheet("Transactions")
    tx.append(["Invoice No", "Narration", "Total"])
    tx.append(["PI-101", "Steel rods purchase", 1180])
    wb.save(xlsx)
    wb.close()

    pred = str(tmp_path / "pred.jsonl")
    run = _run_cli("run", "--input", xlsx, "--out", pred, "--scorer", "keyword")
    assert run.returncode == 0, run.stderr
    rows = [json.loads(line) for line in open(pred, encoding="utf-8") if line.strip()]
    assert len(rows) == 1
    assert rows[0]["invoice_number"] == "PI-101"

    
def test_cli_intake_uses_resilient_workbook_reader(tmp_path):
    source = str(tmp_path / "messy-source.xlsx")
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "FY 2026-27 export"
    tx = wb.create_sheet("Transactions")
    tx.append(["Invoice No", "Narration", "Total"])
    tx.append(["PI-301", "Laptop purchase", 59000])
    wb.save(source)
    wb.close()

    out = str(tmp_path / "intake.xlsx")
    run = _run_cli("intake", "--input", source, "--out", out)
    assert run.returncode == 0, run.stderr

    workbook = load_workbook(out, read_only=True, data_only=True)
    sheet = workbook.active
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    workbook.close()
    assert rows[0][:3] == ["Invoice No", "Narration", "Total"]
    assert rows[1][:3] == ["PI-301", "Laptop purchase", 59000]
    


def test_cli_audit_uses_resilient_workbook_reader(tmp_path):
    source = str(tmp_path / "audit-source.xlsx")
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "VouchPilot export"
    tx = wb.create_sheet("Transactions")
    tx.append(["Invoice No", "voucher_type", "Narration"])
    tx.append(["PI-401", "Purchase", "Steel purchase"])
    wb.save(source)
    wb.close()

    pred = str(tmp_path / "pred.jsonl")
    with open(pred, "w", encoding="utf-8") as handle:
        handle.write(json.dumps({
            "row_id": 1,
            "invoice_number": "PI-401",
            "voucher_type": "Purchase",
            "confidence": 0.95,
        }) + "\n")

    report = str(tmp_path / "audit.json")
    run = _run_cli(
        "audit", "--input", source, "--voucher-col", "voucher_type",
        "--pred", pred, "--report", report,
    )
    assert run.returncode == 0, run.stderr
    result = json.loads(open(report, encoding="utf-8").read())
    assert result["agreement"] == 1.0

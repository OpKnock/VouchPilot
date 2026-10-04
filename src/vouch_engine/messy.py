"""Resilient .xlsx reader for messy Indian accounting exports.

Pre-pass only: normalise.py / baseline.py are frozen, so merged-cell filling,
title-row skipping, Hindi header translation and digit folding live HERE.
Output ``rows`` are raw dicts shaped like ``ingest.read_excel`` rows.
"""

from __future__ import annotations

import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from .i18n import HI_ALIASES, fold_digits, translate_header
from .normalise import ALIASES

_EN_ALIAS_SET: set[str] = {
    a.strip().casefold() for aliases in ALIASES.values() for a in aliases
}


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _materialize_grid(ws: Any) -> tuple[list[list[Any]], int]:
    """Return dense grid with merged cells filled from their top-left value."""
    ranges = list(ws.merged_cells.ranges)
    filled = len(ranges)
    grid: list[list[Any]] = []
    for row in ws.iter_rows():
        out: list[Any] = []
        for cell in row:
            if isinstance(cell, MergedCell):
                val: Any = None
                for rng in ranges:
                    if cell.coordinate in rng:
                        val = ws.cell(
                            row=rng.min_row, column=rng.min_col
                        ).value
                        break
                out.append(val)
            else:
                out.append(cell.value)
        grid.append(out)
    if not grid:
        return [], filled
    # Drop all-empty trailing columns.
    ncols = max(len(r) for r in grid)
    last_used = -1
    for c in range(ncols):
        for r in grid:
            if c < len(r) and not _is_empty(r[c]):
                last_used = c
                break
    if last_used == -1:
        return [], filled
    grid = [r[: last_used + 1] for r in grid]
    # Drop fully-empty trailing rows (keeps footer detection meaningful).
    while grid and all(_is_empty(v) for v in grid[-1]):
        grid.pop()
    return grid, filled


def _score_row(row: list[Any]) -> tuple[float, int]:
    """Score one candidate header row in [0, 1]."""
    non_empty = sum(1 for v in row if not _is_empty(v))
    if non_empty == 0:
        return 0.0, 0
    hits = 0
    for v in row:
        if _is_empty(v):
            continue
        s = str(v).strip()
        if not s:
            continue
        if s.casefold() in _EN_ALIAS_SET or s in HI_ALIASES:
            hits += 1
    return hits / max(1, non_empty), non_empty


def _header_name(raw: Any, idx: int, seen: dict[str, int]) -> str:
    if raw is None:
        base = ""
    elif isinstance(raw, str):
        base = fold_digits(translate_header(raw.strip())).strip()
    else:
        base = fold_digits(str(raw)).strip()
    if base in seen:
        seen[base] += 1
        name = f"{base}_{seen[base]}" if base else f"col_{idx}"
    else:
        seen[base] = 1
        name = base if base else f"col_{idx}"
    return name


def _parse_sheet(ws: Any) -> dict[str, Any]:
    grid, merged = _materialize_grid(ws)
    if not grid:
        return {
            "headers": [],
            "rows": [],
            "profile": {"fill_rate": {}, "n_rows": 0},
            "header_score": 0.0,
            "header_row": 1,
            "decorative_rows": 0,
            "merged_ranges_filled": merged,
            "footer_skipped": False,
        }
    limit = min(15, len(grid))
    scores = [_score_row(grid[i])[0] for i in range(limit)]
    best = max(range(limit), key=lambda i: scores[i])
    header_idx = best if scores[best] >= 0.3 else 0
    header_score = scores[best] if scores[best] >= 0.3 else scores[0]
    header_row_cells = grid[header_idx]
    seen: dict[str, int] = {}
    headers = [
        _header_name(c, i, seen) for i, c in enumerate(header_row_cells)
    ]
    # Data rows below header until a run of 3 fully-empty rows.
    rows: list[dict[str, Any]] = []
    empty_run = 0
    footer_skipped = False
    i = header_idx + 1
    while i < len(grid):
        r = grid[i]
        # Pad short rows so every header has a value.
        if len(r) < len(headers):
            r = list(r) + [None] * (len(headers) - len(r))
        if all(_is_empty(v) for v in r[: len(headers)]):
            empty_run += 1
            if empty_run >= 3:
                rest = grid[i + 1 :]
                if any(
                    not all(_is_empty(v) for v in rr[: len(headers)])
                    for rr in rest
                ):
                    footer_skipped = True
                elif i + 1 < len(grid):
                    # Gap at end with nothing after: still a footer stop.
                    footer_skipped = bool(rest)
                break
            i += 1
            continue
        empty_run = 0
        record: dict[str, Any] = {}
        for j, h in enumerate(headers):
            v = r[j] if j < len(r) else None
            if isinstance(v, str):
                v = fold_digits(v)
            record[h] = v
        rows.append(record)
        i += 1
    n_rows = len(rows)
    fill_rate: dict[str, float] = {}
    for h in headers:
        if n_rows == 0:
            fill_rate[h] = 0.0
        else:
            filled = sum(1 for rec in rows if not _is_empty(rec.get(h)))
            fill_rate[h] = filled / n_rows
    return {
        "headers": headers,
        "rows": rows,
        "profile": {"fill_rate": fill_rate, "n_rows": n_rows},
        "header_score": float(header_score),
        "header_row": header_idx + 1,
        "decorative_rows": header_idx,
        "merged_ranges_filled": merged,
        "footer_skipped": footer_skipped,
    }


def read_messy_xlsx(
    path: str, sheet: str | None = None, pick_best: bool = True
) -> dict[str, Any]:
    """Read a messy workbook into raw rows plus profile and report.

    Args:
        path: .xlsx file path.
        sheet: explicit sheet name; KeyError (naming available sheets) if
            missing. When given, that sheet is used as-is.
        pick_best: when True (default) and ``sheet`` is None, score every
            sheet as ``header_score * (1 + ln(1 + nrows))`` and return the
            best. When False, return the active sheet.

    Returns:
        ``{"headers", "rows", "profile", "report"}`` where ``rows`` are raw
        dicts like ``ingest.read_excel`` rows, ``profile`` holds
        ``{fill_rate, n_rows}`` and ``report`` holds ``{sheet_used,
        sheets_tried, decorative_rows, merged_ranges_filled, header_row,
        footer_skipped}``.
    """
    p = Path(path)
    wb = load_workbook(filename=str(p), data_only=True)
    try:
        available = list(wb.sheetnames)
        if sheet is not None:
            if sheet not in available:
                raise KeyError(
                    f"sheet {sheet!r} not found; available sheets: {available}"
                )
            parsed = _parse_sheet(wb[sheet])
            report = {
                "sheet_used": sheet,
                "sheets_tried": [
                    {
                        "name": sheet,
                        "header_score": parsed["header_score"],
                        "n_rows": parsed["profile"]["n_rows"],
                    }
                ],
                "decorative_rows": parsed["decorative_rows"],
                "merged_ranges_filled": parsed["merged_ranges_filled"],
                "header_row": parsed["header_row"],
                "footer_skipped": parsed["footer_skipped"],
            }
            return {
                "headers": parsed["headers"],
                "rows": parsed["rows"],
                "profile": parsed["profile"],
                "report": report,
            }
        if not pick_best:
            active = wb.active
            name = active.title if active is not None else available[0]
            parsed = _parse_sheet(wb[name])
            report = {
                "sheet_used": name,
                "sheets_tried": [
                    {
                        "name": name,
                        "header_score": parsed["header_score"],
                        "n_rows": parsed["profile"]["n_rows"],
                    }
                ],
                "decorative_rows": parsed["decorative_rows"],
                "merged_ranges_filled": parsed["merged_ranges_filled"],
                "header_row": parsed["header_row"],
                "footer_skipped": parsed["footer_skipped"],
            }
            return {
                "headers": parsed["headers"],
                "rows": parsed["rows"],
                "profile": parsed["profile"],
                "report": report,
            }
        parsed_all: dict[str, dict[str, Any]] = {}
        for name in available:
            parsed_all[name] = _parse_sheet(wb[name])
        best_name = available[0] if available else ""
        best_score = -1.0
        tried: list[dict[str, Any]] = []
        for name in available:
            pa = parsed_all[name]
            nrows = pa["profile"]["n_rows"]
            hs = pa["header_score"]
            sheet_score = hs * (1.0 + math.log1p(nrows))
            tried.append(
                {"name": name, "header_score": hs, "n_rows": nrows}
            )
            if sheet_score > best_score:
                best_score = sheet_score
                best_name = name
        parsed = parsed_all[best_name] if best_name else {
            "headers": [],
            "rows": [],
            "profile": {"fill_rate": {}, "n_rows": 0},
            "header_score": 0.0,
            "header_row": 1,
            "decorative_rows": 0,
            "merged_ranges_filled": 0,
            "footer_skipped": False,
        }
        report = {
            "sheet_used": best_name,
            "sheets_tried": tried,
            "decorative_rows": parsed["decorative_rows"],
            "merged_ranges_filled": parsed["merged_ranges_filled"],
            "header_row": parsed["header_row"],
            "footer_skipped": parsed["footer_skipped"],
        }
        return {
            "headers": parsed["headers"],
            "rows": parsed["rows"],
            "profile": parsed["profile"],
            "report": report,
        }
    finally:
        wb.close()


__all__ = ["read_messy_xlsx"]

"""Excel ingestion with header-row detection and column profiling."""

from __future__ import annotations

from pathlib import Path
from typing import Any


def _is_non_empty(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str) and value.strip() == "":
        return False
    return True


def _is_numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _is_text(value: Any) -> bool:
    return isinstance(value, str) and value.strip() != ""


def detect_header_row(rows: list[list[Any]], scan_limit: int = 10) -> int:
    """Return 0-based index of the header row within *rows*.

    Rule: first row (within first ``scan_limit`` rows) with >=3 non-empty
    cells and more text cells than numeric cells wins.
    Fallback: first non-empty row, else 0.
    """
    limit = min(scan_limit, len(rows))
    for idx in range(limit):
        row = rows[idx]
        non_empty = sum(1 for v in row if _is_non_empty(v))
        if non_empty < 3:
            continue
        n_text = sum(1 for v in row if _is_text(v))
        n_num = sum(1 for v in row if _is_numeric(v))
        if n_text > n_num:
            return idx
    for idx, row in enumerate(rows):
        if any(_is_non_empty(v) for v in row):
            return idx
    return 0


def read_excel(path: str | Path) -> dict[str, Any]:
    """Read an .xlsx file into raw rows plus a column profile.

    Returns ``{"headers", "rows", "profile"}`` where ``rows`` is a list of
    raw dicts (header -> cell value) and ``profile`` holds
    ``{fill_rate, n_rows, header_row}``. ``header_row`` is 1-based Excel
    row number.
    """
    from openpyxl import load_workbook

    p = Path(path)
    wb = load_workbook(filename=str(p), read_only=True, data_only=True)
    ws = wb.active
    assert ws is not None
    all_rows: list[list[Any]] = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    if not all_rows:
        return {"headers": [], "rows": [], "profile": {"fill_rate": {}, "n_rows": 0, "header_row": 1}}

    header_idx = detect_header_row(all_rows)
    raw_headers = all_rows[header_idx]
    headers: list[str] = []
    seen: dict[str, int] = {}
    for cell in raw_headers:
        name = "" if cell is None else str(cell).strip()
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}" if name else f"col_{len(headers)}"
        else:
            seen[name] = 1
            if name == "":
                name = f"col_{len(headers)}"
        headers.append(name)

    data_rows = all_rows[header_idx + 1 :]
    rows: list[dict[str, Any]] = []
    for r in data_rows:
        if not any(_is_non_empty(v) for v in r):
            continue
        record: dict[str, Any] = {}
        for i, h in enumerate(headers):
            record[h] = r[i] if i < len(r) else None
        rows.append(record)

    n_rows = len(rows)
    fill_rate: dict[str, float] = {}
    for h in headers:
        if n_rows == 0:
            fill_rate[h] = 0.0
        else:
            filled = sum(1 for rec in rows if _is_non_empty(rec.get(h)))
            fill_rate[h] = filled / n_rows

    return {
        "headers": headers,
        "rows": rows,
        "profile": {"fill_rate": fill_rate, "n_rows": n_rows, "header_row": header_idx + 1},
    }

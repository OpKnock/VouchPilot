"""VouchEngine CLI (Team-C): run / evaluate / gold / doctor.

``run`` wires the Team-A deterministic pipeline (ingest -> normalise ->
perspective -> evidence), the Team-B scorer, and Team-A validate writers by
*importing* those modules against their contracts:

- ``ingest.read_excel(path)`` -> ``{headers, rows, profile}``
- ``normalise.map_columns(headers, sample_rows)`` -> ``{raw: canon | None}``
- ``normalise.to_canonical(rows, mapping)`` -> ``[CanonicalRow]``
- ``perspective.resolve(rows)`` -> ``(self_entity, perspectives)`` (in place)
- ``evidence.extract(row)`` -> ``(tags, mask)``
- ``scorer.StubScorer / LlamaServerScorer`` with
  ``.predict(row, tags, mask)`` -> ``(label, conf, top_k)``
- ``validate.validate_prediction / write_jsonl / write_xlsx``

If an upstream module is missing (or the full pipeline raises), ``run``
degrades to a minimal stdlib+openpyxl path that still emits one valid
:class:`Prediction` per row and reports the fallback on stderr, so rows are
never silently dropped.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import importlib
import json
import math
import os
import sys
import time
import urllib.request

from .labels import LABEL_NAMES
from .schemas import Prediction

UPSTREAM_MODULES = (
    "ingest",
    "normalise",
    "perspective",
    "evidence",
    "baseline",
    "scorer",
    "validate",
)
OWN_MODULES = ("gold", "evaluate", "labels", "schemas", "retriever", "calibrate", "challenger")

INVOICE_KEYS = ("Bill No", "Invoice No", "Doc No", "Invoice Number", "BillNo", "invoice_number")
FALLBACK_LABEL = "Other / Miscellaneous"


def _load(mod_name: str):
    try:
        return importlib.import_module(f".{mod_name}", package=__package__)
    except Exception:
        return None


def _extract_invoice(raw: dict, row_id: int) -> str:
    for key in INVOICE_KEYS:
        value = (raw or {}).get(key)
        if value is not None and str(value).strip() != "":
            return str(value).strip()
    return f"ROW-{row_id}"


def _read_xlsx_fallback(path: str) -> tuple:
    """Minimal reader mirroring the ingest contract (headers + raw rows)."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    raw_rows = []
    for values in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        raw_rows.append({h: (v if v is not None else "") for h, v in zip(headers, values) if h})
    return headers, raw_rows


class _KeywordAdapter:
    """A0 measurement-only scorer wrapping baseline.keyword_predict."""

    def __init__(self, baseline_mod):
        self._baseline = baseline_mod

    def predict(self, row, tags, mask, exemplars=None):
        _ = exemplars
        label, conf = self._baseline.keyword_predict(row, tags)
        return label, conf, [[label, 1.0]]


def _build_scorer(scorer_mod, baseline_mod, which: str, endpoint: str):
    if which == "server":
        return scorer_mod.LlamaServerScorer(endpoint=endpoint)
    if which == "keyword":
        return _KeywordAdapter(baseline_mod)
    if which == "vouchpilot":
        try:
            from extensions.vouchpilot_scorer import VouchPilotScorer

            return VouchPilotScorer(base="keyword", endpoint=endpoint)
        except Exception as exc:
            print(
                f"WARN: vouchpilot scorer unavailable ({exc}); keyword fallback",
                file=sys.stderr,
            )
            return _KeywordAdapter(baseline_mod)
    return scorer_mod.StubScorer()


def _as_prediction(validate_mod, rec: dict):
    if validate_mod is not None and hasattr(validate_mod, "validate_prediction"):
        return validate_mod.validate_prediction(rec)
    return Prediction.model_validate(rec)


def _write_outputs(validate_mod, preds: list, out_path: str, xlsx_path) -> None:
    if validate_mod is not None and hasattr(validate_mod, "write_jsonl"):
        validate_mod.write_jsonl(preds, out_path)
    else:
        _write_jsonl_local(preds, out_path)
    if xlsx_path:
        if validate_mod is not None and hasattr(validate_mod, "write_xlsx"):
            validate_mod.write_xlsx(preds, xlsx_path)
        else:
            _write_xlsx_local(preds, xlsx_path)


def _write_jsonl_local(preds: list, path: str) -> None:
    directory = os.path.dirname(os.path.abspath(path))
    os.makedirs(directory, exist_ok=True)
    with open(path, "w", encoding="utf-8") as handle:
        for pred in preds:
            data = pred.model_dump() if isinstance(pred, Prediction) else pred
            handle.write(json.dumps(data) + "\n")


def _write_xlsx_local(preds: list, path: str) -> None:
    from openpyxl import Workbook

    directory = os.path.dirname(os.path.abspath(path))
    os.makedirs(directory, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Predictions"
    sheet.append(["row_id", "invoice_number", "voucher_type", "confidence", "needs_review"])
    for pred in preds:
        data = pred.model_dump() if isinstance(pred, Prediction) else pred
        sheet.append(
            [
                data["row_id"],
                data["invoice_number"],
                data["voucher_type"],
                data["confidence"],
                data["needs_review"],
            ]
        )
    workbook.save(path)


def _xlsx_out_path(out_path: str, xlsx_arg) -> str | None:
    if xlsx_arg is None:
        return None
    if xlsx_arg == "auto":
        base, _ = os.path.splitext(out_path)
        return base + ".xlsx"
    return str(xlsx_arg)


def _coerce_record(row_id: int, invoice: str, label, conf, top_k, tags) -> dict:
    try:
        confidence = min(1.0, max(0.0, float(conf)))
    except (TypeError, ValueError):
        confidence = 0.0
    if not isinstance(top_k, list) or not top_k:
        top_k = [[label, 1.0]]
    return {
        "row_id": row_id,
        "invoice_number": invoice or f"ROW-{row_id}",
        "voucher_type": label,
        "confidence": confidence,
        "needs_review": bool(confidence < 0.5),
        "top_k": top_k,
        "evidence": [str(t) for t in tags] if isinstance(tags, list) else [],
    }


def _finalise(
    preds: list, invalid: int, out_path: str, xlsx_path, validate_mod, writer: str
) -> int:
    _write_outputs(validate_mod, preds, out_path, xlsx_path)
    needs_review = sum(1 for p in preds if p.needs_review)
    print(f"n_rows={len(preds)} invalid={invalid} needs_review={needs_review} writer={writer}")
    return 0


def _load_exemplar_index(prefix: str):
    """Load a retriever index from PREFIX or PREFIX_exemplars.json."""
    import json as _json

    candidates = [prefix, prefix + "_exemplars.json"]
    for cand in candidates:
        if cand and os.path.exists(cand) and os.path.isfile(cand):
            with open(cand, encoding="utf-8") as handle:
                data = _json.load(handle)
            if isinstance(data, dict) and "exemplars" in data:
                return data
            if isinstance(data, list):
                return {"exemplars": data}
            return {"exemplars": []}
    raise FileNotFoundError(f"exemplar index not found: {prefix}")


def _margin_of(top_k) -> float:
    try:
        if isinstance(top_k, list) and len(top_k) >= 2:
            return float(top_k[0][1]) - float(top_k[1][1])
    except (TypeError, ValueError, IndexError):
        pass
    return 1.0


def _row_desc_for_challenger(crow: dict) -> str:
    narr = crow.get("narration", "") if isinstance(crow, dict) else ""
    items = crow.get("items", []) if isinstance(crow, dict) else []
    return f"narration={narr} items={items}"


def _predict_with_exemplars(scorer_obj, crow, tags, mask, exemplars):
    try:
        return scorer_obj.predict(crow, tags, mask, exemplars=exemplars)
    except TypeError:
        return scorer_obj.predict(crow, tags, mask)


def _first_pass_state(ctx: dict, idx: int) -> dict:
    """Score one row (thread-safe: all callees are stateless)."""
    mods, scorer_obj = ctx["mods"], ctx["scorer_obj"]
    canonical, raw_rows = ctx["canonical"], ctx["raw_rows"]
    pos = idx + 1
    crow = canonical[idx]
    row_id = crow.get("row_id", pos)
    raw = raw_rows[idx] if idx < len(raw_rows) else {}
    invoice = (crow.get("doc") or {}).get("invoice_number")
    invoice = invoice or _extract_invoice(raw, row_id)
    try:
        tags, mask = mods["evidence"].extract(crow)
        row_ex: list = []
        retrieve, index = ctx["retrieve"], ctx["index"]
        if index is not None and retrieve is not None:
            try:
                row_ex = retrieve(
                    index, tags, k=ctx["k"], per_class_cap=2, exclude_row_id=row_id
                )
            except TypeError:
                row_ex = retrieve(index, tags, k=ctx["k"])
        label, conf, top_k = _predict_with_exemplars(scorer_obj, crow, tags, mask, row_ex)
        return {
            "ok": True, "crow": crow, "row_id": row_id, "invoice": invoice,
            "tags": list(tags), "label": label, "conf": conf, "top_k": top_k,
        }
    except Exception as exc:
        print(f"WARN: row {row_id} pipeline failed ({exc}); coerced", file=sys.stderr)
        return {"ok": False, "row_id": row_id, "invoice": invoice}


def _challenge_state(ctx: dict, n: int) -> tuple:
    """Re-check one row (thread-safe); returns (n, result-tuple or error str)."""
    states = ctx["states"]
    s = states[n]
    label_a = s["top_k"][0][0]
    label_b = s["top_k"][1][0]
    complete_fn = None
    if ctx["escalation_scorer"] is not None:
        complete_fn = ctx["escalation_scorer"]._complete
    elif hasattr(ctx["scorer_obj"], "_complete"):
        complete_fn = ctx["scorer_obj"]._complete
    if complete_fn is None:
        return (n, None, "no-complete-fn")
    try:
        from .challenger import recheck as _recheck

        desc = _row_desc_for_challenger(s["crow"])
        winner, chal_conf = _recheck(desc, s["tags"], label_a, label_b, complete_fn)
        return (n, (winner, float(chal_conf)), None)
    except Exception as exc:
        return (n, None, str(exc))


def _cmd_run_full(args, mods) -> int:
    data = mods["ingest"].read_excel(args.input)
    headers, raw_rows = list(data.get("headers", [])), list(data.get("rows", []))
    mapping = mods["normalise"].map_columns(headers, raw_rows)
    canonical = mods["normalise"].to_canonical(raw_rows, mapping)
    mods["perspective"].resolve(canonical)
    scorer_obj = _build_scorer(mods["scorer"], mods["baseline"], args.scorer, args.endpoint)
    exemplars_prefix = getattr(args, "exemplars", None)
    top_k_n = int(getattr(args, "k", 5) or 5)
    margin_thr = float(getattr(args, "margin", 0.15))
    escalation_url = getattr(args, "escalation_endpoint", None)
    calibrator_path = getattr(args, "calibrator", None)

    index = None
    if exemplars_prefix:
        try:
            from .retriever import retrieve as _retrieve

            index = _load_exemplar_index(exemplars_prefix)
        except Exception as exc:
            print(f"WARN: exemplar index load failed ({exc}); continuing", file=sys.stderr)
            index = None
            _retrieve = None  # type: ignore[assignment]
    else:
        _retrieve = None  # type: ignore[assignment]

    cal_t = None
    if calibrator_path:
        try:
            from .calibrate import load_calibrator as _load_cal

            cal_t = float(_load_cal(calibrator_path).get("T", 1.0))
        except Exception as exc:
            print(f"WARN: calibrator load failed ({exc}); continuing", file=sys.stderr)
            cal_t = None

    escalation_scorer = None
    if escalation_url:
        try:
            escalation_scorer = mods["scorer"].LlamaServerScorer(endpoint=escalation_url)
        except Exception as exc:
            print(f"WARN: escalation scorer init failed ({exc})", file=sys.stderr)
            escalation_scorer = None

    max_rate = float(getattr(args, "max_challenge_rate", 1.0))
    preds, invalid = [], 0
    challenged = 0
    challenger_skipped = 0
    total_rows = len(canonical)
    start = max(0, getattr(args, "offset", 0) or 0)
    limit = getattr(args, "limit", 0) or 0
    stop = start + limit if limit else total_rows
    indices = list(range(start, min(stop, total_rows)))
    workers = max(1, int(getattr(args, "workers", 4) or 1))
    # Pass 1: score every row, record margins (no challenger yet).
    # Order-preserving: executor.map yields in input order.
    ctx = {
        "mods": mods, "scorer_obj": scorer_obj, "index": index,
        "retrieve": _retrieve, "k": top_k_n,
        "canonical": canonical, "raw_rows": raw_rows,
    }
    states: list[dict] = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as pool:
        for out_pos, state in enumerate(
            pool.map(lambda i: _first_pass_state(ctx, i), indices), start=1
        ):
            states.append(state)
            if out_pos % 10 == 0 or out_pos == len(indices):
                print(f"progress {out_pos}/{len(indices)}", file=sys.stderr)
    # Pass 2: challenge the lowest-margin rows within budget (rate cap).
    budget = math.ceil(max_rate * len(states)) if max_rate > 0 else 0
    ranked = sorted(
        (
            (_margin_of(s["top_k"]), n)
            for n, s in enumerate(states)
            if s["ok"] and isinstance(s["top_k"], list) and len(s["top_k"]) >= 2
            and _margin_of(s["top_k"]) < margin_thr
        ),
        key=lambda t: (t[0], t[1]),
    )
    chosen = {n for _, n in ranked[:budget]}
    ctx["states"] = states
    ctx["escalation_scorer"] = escalation_scorer
    with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as pool:
        for n, result, err in pool.map(
            lambda i: _challenge_state(ctx, i), sorted(chosen)
        ):
            s = states[n]
            if err == "no-complete-fn":
                challenger_skipped += 1
                s["force_review"] = True
                continue
            if err is not None:
                print(f"WARN: row {s['row_id']} challenger failed ({err})", file=sys.stderr)
                challenger_skipped += 1
                s["force_review"] = True
                continue
            winner, chal_conf = result
            s["label"], s["conf"] = winner, float(chal_conf)
            if s["top_k"] and s["top_k"][0][0] == winner:
                s["top_k"][0][1] = float(chal_conf)
            s["tags"] = s["tags"] + ["CHALLENGED"]
            challenged += 1
    # Pass 3: calibrate, coerce, validate in original order.
    for s in states:
        if not s["ok"]:
            rec = _coerce_record(
                s["row_id"], s["invoice"], FALLBACK_LABEL, 0.0,
                [[FALLBACK_LABEL, 1.0]], ["pipeline-failed"],
            )
            invalid += 1
        else:
            label, conf, top_k = s["label"], s["conf"], s["top_k"]
            out_tags = s["tags"]
            force_review = bool(s.get("force_review", False))
            if cal_t is not None:
                try:
                    from .calibrate import apply_conf as _apply2

                    conf = float(_apply2(float(conf), cal_t))
                    if top_k and top_k[0][0] == label:
                        top_k[0][1] = float(conf)
                except (TypeError, ValueError):
                    pass
            rec = _coerce_record(s["row_id"], s["invoice"], label, conf, top_k, out_tags)
            if force_review:
                rec["needs_review"] = True
        try:
            preds.append(_as_prediction(mods["validate"], rec))
        except Exception as exc:
            print(f"WARN: row {s['row_id']} invalid ({exc}); coerced", file=sys.stderr)
            invalid += 1
            rec.update(
                voucher_type=FALLBACK_LABEL,
                confidence=0.0,
                needs_review=True,
                top_k=[[FALLBACK_LABEL, 1.0]],
            )
            preds.append(_as_prediction(mods["validate"], rec))
    code = _finalise(
        preds, invalid, args.out, _xlsx_out_path(args.out, args.xlsx), mods["validate"], "validate"
    )
    if exemplars_prefix or escalation_url or margin_thr != 0.15:
        print(
            f"challenged={challenged} challenger_skipped={challenger_skipped} "
            f"budget_rate={max_rate}",
            file=sys.stderr,
        )
    return code


def _cmd_run_fallback(args, validate_mod, reason: str) -> int:
    print(f"WARN: {reason}; using fallback path", file=sys.stderr)
    _, raw_rows = _read_xlsx_fallback(args.input)
    preds, invalid = [], 0
    for pos, raw in enumerate(raw_rows, start=1):
        label = LABEL_NAMES[(pos - 1) % len(LABEL_NAMES)]
        rec = _coerce_record(
            pos, _extract_invoice(raw, pos), label, 0.5, [[label, 1.0]], ["fallback"]
        )
        try:
            preds.append(_as_prediction(validate_mod, rec))
        except Exception:
            invalid += 1
    return _finalise(
        preds, invalid, args.out, _xlsx_out_path(args.out, args.xlsx), validate_mod, "fallback"
    )


def cmd_run(args) -> int:
    if not os.path.exists(args.input):
        print(f"ERROR: input not found: {args.input}", file=sys.stderr)
        return 2
    mods = {name: _load(name) for name in UPSTREAM_MODULES}
    missing = [name for name, mod in mods.items() if mod is None]
    if missing:
        return _cmd_run_fallback(
            args, mods["validate"], f"upstream modules missing: {', '.join(missing)}"
        )
    try:
        return _cmd_run_full(args, mods)
    except Exception as exc:
        return _cmd_run_fallback(args, mods["validate"], f"full pipeline failed ({exc})")


def _load_json(path: str):
    with open(path, encoding="utf-8") as handle:
        return json.load(handle)


def _load_pred_list(path: str) -> list:
    preds = []
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if line:
                preds.append(json.loads(line))
    return preds


def cmd_evaluate(args) -> int:
    from .evaluate import compute_metrics, write_report

    gold = _load_json(args.gold)
    if isinstance(gold, dict):
        gold = gold.get("labels", gold.get("rows", []))
    pred_list = _load_pred_list(args.pred)
    by_id = {}
    conf_by_id: dict[int, float] = {}
    for entry in pred_list:
        rid = entry.get("row_id")
        if rid is not None:
            try:
                by_id[int(rid)] = entry.get("voucher_type")
                conf_by_id[int(rid)] = float(entry.get("confidence", 0.5))
            except (TypeError, ValueError):
                continue
    y_true, y_pred, y_conf = [], [], []
    for pos, entry in enumerate(gold, start=1):
        rid = entry.get("row_id", pos)
        y_true.append(entry.get("voucher_type", entry.get("label")))
        hit = by_id.get(int(rid)) if isinstance(rid, int) else None
        if hit is None and pos - 1 < len(pred_list):
            hit = pred_list[pos - 1].get("voucher_type")
        y_pred.append(hit if hit in LABEL_NAMES else FALLBACK_LABEL)
        if isinstance(rid, int) and rid in conf_by_id:
            y_conf.append(conf_by_id[rid])
        elif pos - 1 < len(pred_list):
            try:
                y_conf.append(float(pred_list[pos - 1].get("confidence", 0.5)))
            except (TypeError, ValueError):
                y_conf.append(0.5)
        else:
            y_conf.append(0.5)
    report = compute_metrics(y_true, y_pred)
    calibrator_path = getattr(args, "calibrator", None)
    if calibrator_path:
        from .calibrate import apply_conf, coverage_accuracy, ece
        from .calibrate import load_calibrator
        from .challenger import PAIRS, pairwise_f1

        try:
            cal = load_calibrator(calibrator_path)
            temp = float(cal.get("T", 1.0))
        except Exception as exc:
            print(f"WARN: calibrator load failed ({exc})", file=sys.stderr)
            temp = 1.0
        corrects = [bool(t == p) for t, p in zip(y_true, y_pred)]
        ece_before = float(ece(y_conf, corrects))
        cal_confs = [float(apply_conf(c, temp)) for c in y_conf]
        ece_after = float(ece(cal_confs, corrects))
        coverage = coverage_accuracy(cal_confs, corrects, steps=10)
        pair_scores = pairwise_f1(y_true, y_pred, PAIRS)
        pair_str = {f"{a} vs {b}": float(v) for (a, b), v in pair_scores.items()}
        report["ece"] = ece_after
        report["ece_before"] = ece_before
        report["ece_after"] = ece_after
        report["coverage"] = coverage
        report["coverage_accuracy"] = coverage
        report["pairwise_f1"] = pair_str
        report["calibrator"] = calibrator_path
        report["temperature"] = temp
    write_report(report, args.report)
    print(
        f"accuracy={report['accuracy']:.4f} macro_f1={report['macro_f1']:.4f} "
        f"n_rows={report['n_rows']}"
    )
    return 0


def cmd_calibrate(args) -> int:
    from .calibrate import apply_conf, ece, fit_temperature, save_calibrator

    gold = _load_json(args.gold)
    if isinstance(gold, dict):
        gold = gold.get("labels", gold.get("rows", []))
    pred_list = _load_pred_list(args.pred)
    by_id: dict[int, dict] = {}
    for entry in pred_list:
        rid = entry.get("row_id")
        if rid is not None:
            try:
                by_id[int(rid)] = entry
            except (TypeError, ValueError):
                continue
    confs, corrects, items = [], [], []
    for pos, entry in enumerate(gold, start=1):
        rid = entry.get("row_id", pos)
        true_label = entry.get("voucher_type", entry.get("label"))
        pred = by_id.get(int(rid)) if isinstance(rid, int) else None
        if pred is None and pos - 1 < len(pred_list):
            pred = pred_list[pos - 1]
        if not isinstance(pred, dict):
            continue
        try:
            conf = float(pred.get("confidence", 0.5))
        except (TypeError, ValueError):
            conf = 0.5
        ok = bool(pred.get("voucher_type") == true_label)
        confs.append(conf)
        corrects.append(ok)
        items.append({"conf": conf, "correct": ok})
    temp = float(fit_temperature(items))
    ece_before = float(ece(confs, corrects))
    cal_confs = [float(apply_conf(c, temp)) for c in confs]
    ece_after = float(ece(cal_confs, corrects))
    save_calibrator(args.out, temp, len(items), ece_before, ece_after)
    print(f"T={temp:.4f} n={len(items)} ece_before={ece_before:.4f} ece_after={ece_after:.4f}")
    return 0


def cmd_gold(args) -> int:
    from .gold import build_dataset

    result = build_dataset(args.n, args.seed, args.out)
    print(f"wrote {result['xlsx']} + {result['labels']} n_rows={result['n_rows']}")
    return 0


def cmd_doctor(args) -> int:
    ok = True
    for name in (*UPSTREAM_MODULES, *OWN_MODULES):
        mod = _load(name)
        print(f"import vouch_engine.{name}: {'OK' if mod is not None else 'FAIL'}")
        ok = ok and mod is not None

    weights_hint = getattr(args, "weights", None)
    candidates = []
    if weights_hint:
        candidates.append(weights_hint)
    if os.environ.get("VOUCH_MODEL_PATH"):
        candidates.append(os.environ["VOUCH_MODEL_PATH"])
    candidates.extend(["models", "weights"])
    found = None
    for cand in candidates:
        if os.path.isfile(cand) and cand.endswith(".gguf"):
            found = cand
            break
        if os.path.isdir(cand):
            hits = [f for f in os.listdir(cand) if f.endswith(".gguf")]
            if hits:
                found = os.path.join(cand, hits[0])
                break
    if found:
        print(f"weights dir check: OK ({found})")
    else:
        print(f"weights dir check: FAIL (no .gguf under {candidates})")
        ok = False

    endpoint = (getattr(args, "endpoint", None) or "http://localhost:8080").rstrip("/")
    try:
        with urllib.request.urlopen(endpoint + "/health", timeout=2) as resp:
            healthy = 200 <= resp.status < 300
    except Exception as exc:
        print(f"server /health check: FAIL ({exc})")
        ok = False
    else:
        print(f"server /health check: {'OK' if healthy else 'FAIL'} ({endpoint})")
        ok = ok and healthy
    return 0 if ok else 1


def _agent_sidecars(out_path: str) -> dict:
    stem, _ = os.path.splitext(os.path.abspath(out_path))
    return {
        "approvals": stem + ".approvals.json",
        "trace": stem + ".trace.jsonl",
        "inspection": stem + ".inspection.json",
    }


def _agent_setup(args, mods: dict) -> dict:
    """Mirror of the run-pipeline setup: scorer, exemplars, escalation, calibrator."""
    scorer_obj = _build_scorer(mods["scorer"], mods["baseline"], args.scorer, args.endpoint)
    index, retrieve = None, None
    if getattr(args, "exemplars", None):
        try:
            from .retriever import retrieve as _retrieve

            index = _load_exemplar_index(args.exemplars)
            retrieve = _retrieve
        except Exception as exc:
            print(f"WARN: exemplar index load failed ({exc}); continuing", file=sys.stderr)
    cal_t = None
    if getattr(args, "calibrator", None):
        try:
            from .calibrate import load_calibrator as _load_cal

            cal_t = float(_load_cal(args.calibrator).get("T", 1.0))
        except Exception as exc:
            print(f"WARN: calibrator load failed ({exc}); continuing", file=sys.stderr)
    escalation_scorer = None
    if getattr(args, "escalation_endpoint", None):
        try:
            escalation_scorer = mods["scorer"].LlamaServerScorer(
                endpoint=args.escalation_endpoint
            )
        except Exception as exc:
            print(f"WARN: escalation scorer init failed ({exc})", file=sys.stderr)
    return {
        "mods": mods, "scorer_obj": scorer_obj, "index": index, "retrieve": retrieve,
        "k": int(getattr(args, "k", 5) or 5), "escalation_scorer": escalation_scorer,
        "cal_t": cal_t, "workers": max(1, int(getattr(args, "workers", 1) or 1)),
    }


def _agent_finalize(states: list, args, mods: dict, cal_t) -> tuple[list, int]:
    preds, invalid = [], 0
    for s in states:
        if not s["ok"]:
            rec = _coerce_record(
                s["row_id"], s["invoice"], FALLBACK_LABEL, 0.0,
                [[FALLBACK_LABEL, 1.0]], ["pipeline-failed"],
            )
            invalid += 1
        else:
            label, conf, top_k = s["label"], s["conf"], s["top_k"]
            if cal_t is not None:
                try:
                    from .calibrate import apply_conf as _apply2

                    conf = float(_apply2(float(conf), cal_t))
                    if top_k and top_k[0][0] == label:
                        top_k[0][1] = float(conf)
                except (TypeError, ValueError):
                    pass
            rec = _coerce_record(s["row_id"], s["invoice"], label, conf, top_k, s["tags"])
            if s.get("force_review"):
                rec["needs_review"] = True
            elif s.get("cleared"):
                rec["needs_review"] = False
        try:
            preds.append(_as_prediction(mods["validate"], rec))
        except Exception as exc:
            print(f"WARN: row {s['row_id']} invalid ({exc}); coerced", file=sys.stderr)
            invalid += 1
            rec.update(
                voucher_type=FALLBACK_LABEL, confidence=0.0, needs_review=True,
                top_k=[[FALLBACK_LABEL, 1.0]],
            )
            preds.append(_as_prediction(mods["validate"], rec))
    return preds, invalid


def _parse_gate(gate: str) -> tuple[str, str | None]:
    if gate == "auto":
        return "auto", None
    if gate == "cli":
        return "cli", None
    if gate.startswith("file:"):
        return "file", gate[5:]
    raise ValueError("gate must be auto | cli | file:PATH")


def cmd_agent_run(args) -> int:
    from . import agent as agent_mod

    if not os.path.exists(args.input):
        print(f"ERROR: input not found: {args.input}", file=sys.stderr)
        return 2
    mods = {name: _load(name) for name in UPSTREAM_MODULES}
    missing = [name for name, mod in mods.items() if mod is None]
    if missing:
        print(f"ERROR: modules missing: {', '.join(missing)}", file=sys.stderr)
        return 2
    gate, gate_path = _parse_gate(args.gate)
    sidecars = _agent_sidecars(args.out)
    state = agent_mod.AgentState(os.path.dirname(os.path.abspath(args.out)))
    inspected = agent_mod.inspect_sheet(mods, args.input, state)
    rep = inspected["report"]
    print(
        f"inspect: n_rows={rep['n_rows']} mapped={rep['mapped_columns']}/{rep['total_columns']} "
        f"self={rep['self_entity']} perspectives={rep['perspectives']} flags={rep['red_flags']}"
    )
    ctx = _agent_setup(args, mods)
    ctx.update(
        {"canonical": inspected["canonical"], "raw_rows": inspected["raw_rows"]}
    )
    total = len(inspected["canonical"])
    start = max(0, args.offset or 0)
    stop = start + args.limit if args.limit else total
    indices = list(range(start, min(stop, total)))
    states = agent_mod.classify_all(ctx, indices, ctx["workers"], state)
    challenged, skipped = agent_mod.challenge_budgeted(
        ctx, states, float(args.margin), float(args.max_challenge_rate), state
    )
    queue = agent_mod.build_review_queue(states)
    print(f"classify: n={len(states)} challenged={challenged} skipped={skipped} "
          f"review_queue={len(queue)}")
    if args.pending_out:
        with open(args.pending_out, "w", encoding="utf-8") as handle:
            json.dump({"source": args.input, "queue": queue}, handle, indent=1)
    if gate == "auto":
        decisions = []
    elif gate == "cli":
        decisions = agent_mod.gate_cli(queue)
    else:
        assert gate_path is not None
        decisions = agent_mod.gate_file(gate_path)
    approvals = agent_mod.apply_gate(states, decisions, gate, state)
    preds, invalid = _agent_finalize(states, args, mods, ctx["cal_t"])
    _write_outputs(mods["validate"], preds, args.out,
                   _xlsx_out_path(args.out, args.xlsx))
    with open(sidecars["approvals"], "w", encoding="utf-8") as handle:
        json.dump({"approvals": approvals}, handle, indent=1)
    with open(sidecars["inspection"], "w", encoding="utf-8") as handle:
        json.dump(rep, handle, indent=1)
    state.write_trace(sidecars["trace"])
    needs = sum(1 for p in preds if p.needs_review)
    print(f"export: n_rows={len(preds)} invalid={invalid} needs_review={needs} "
          f"approvals={len(approvals)} trace={state.tool_calls} calls")
    return 0


def cmd_agent_template(args) -> int:
    from . import agent as agent_mod

    if not os.path.exists(args.input):
        print(f"ERROR: input not found: {args.input}", file=sys.stderr)
        return 2
    mods = {name: _load(name) for name in UPSTREAM_MODULES}
    missing = [name for name, mod in mods.items() if mod is None]
    if missing:
        print(f"ERROR: modules missing: {', '.join(missing)}", file=sys.stderr)
        return 2
    state = agent_mod.AgentState(os.path.dirname(os.path.abspath(args.out)))
    inspected = agent_mod.inspect_sheet(mods, args.input, state)
    ctx = _agent_setup(args, mods)
    ctx.update(
        {"canonical": inspected["canonical"], "raw_rows": inspected["raw_rows"]}
    )
    indices = list(range(len(inspected["canonical"])))
    states = agent_mod.classify_all(ctx, indices, ctx["workers"], state)
    agent_mod.challenge_budgeted(
        ctx, states, float(args.margin), float(args.max_challenge_rate), state
    )
    queue = agent_mod.build_review_queue(states)
    with open(args.out, "w", encoding="utf-8") as handle:
        json.dump({"source": args.input, "queue": queue}, handle, indent=1)
    print(f"pending: {len(queue)} rows need review -> {args.out}")
    return 0


def cmd_audit(args) -> int:
    from . import audit as audit_mod

    if not os.path.exists(args.input):
        print(f"ERROR: input not found: {args.input}", file=sys.stderr)
        return 2
    mods = {name: _load(name) for name in UPSTREAM_MODULES}
    if mods.get("ingest") is None:
        print("ERROR: ingest module missing", file=sys.stderr)
        return 2
    data = mods["ingest"].read_excel(args.input)
    records = []
    for pos, raw in enumerate(data.get("rows", []), start=1):
        voucher = raw.get(args.voucher_col, "")
        if str(voucher).strip() == "":
            continue
        records.append(
            {
                "row_id": pos,
                "invoice_number": str(raw.get("invoice_number", f"ROW-{pos}")),
                "voucher_type": str(voucher).strip(),
            }
        )
    if not records:
        print(f"ERROR: no recorded '{args.voucher_col}' values found", file=sys.stderr)
        return 2
    try:
        with open(args.pred, encoding="utf-8") as handle:
            predictions = [json.loads(line) for line in handle if line.strip()]
    except OSError as exc:
        print(f"ERROR: cannot read predictions ({exc})", file=sys.stderr)
        return 2
    report = audit_mod.audit_recorded_vs_predicted(records, predictions)
    with open(args.report, "w", encoding="utf-8") as handle:
        json.dump(report, handle, indent=1)
    print(f"agreement={report['agreement']:.4f} n_rows={report['n_rows']} "
          f"disagreements={len(report['disagreements'])}")
    return 0


def cmd_robust(args) -> int:
    from . import robust as robust_mod

    try:
        with open(args.gold + "_labels.json", encoding="utf-8") as handle:
            labels = json.load(handle)
    except OSError as exc:
        print(f"ERROR: cannot read gold labels ({exc})", file=sys.stderr)
        return 2
    rates = [float(x) for x in str(args.rates).split(",") if x.strip() != ""]
    rename_rates = [float(x) for x in str(args.rename_rates).split(",") if x.strip() != ""]
    report = robust_mod.sweep(
        args.gold + ".xlsx", labels, rates, rename_rates, args.scorer, args.seed
    )
    with open(args.out, "w", encoding="utf-8") as handle:
        json.dump(report, handle, indent=1)
    print(f"base_macro_f1={report['base_macro_f1']:.4f} "
          f"o4_within_10pts={report['verdict_o4_drop_within_10pts']} "
          f"o5_within_5pts={report['verdict_o5_rename_within_5pts']}")
    return 0


def cmd_intake(args) -> int:
    from . import intake as intake_mod

    if not os.path.exists(args.input):
        print(f"ERROR: input not found: {args.input}", file=sys.stderr)
        return 2
    langs = tuple(dict.fromkeys(
        [part.strip() for part in str(args.langs).split(",") if part.strip()]
    )) or ("hin", "eng")
    try:
        report = intake_mod.intake_to_xlsx(args.input, args.out, langs)
    except intake_mod.IntakeError as exc:
        print(f"ERROR: intake failed ({exc})", file=sys.stderr)
        return 2
    if report.get("ocr_used"):
        print("WARN: OCR was used; verify extracted fields before classifying",
              file=sys.stderr)
    print(f"kind={report.get('kind')} pages={report.get('pages')} "
          f"ocr_used={report.get('ocr_used')} rows={report.get('rows')} "
          f"-> {args.out}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="python -m vouch_engine")
    sub = parser.add_subparsers(dest="command", required=True)

    run_p = sub.add_parser("run", help="Classify an .xlsx sheet to predictions.jsonl")
    run_p.add_argument("--input", required=True)
    run_p.add_argument("--out", required=True)
    run_p.add_argument(
        "--xlsx",
        nargs="?",
        const="auto",
        default=None,
        help="Also write predictions xlsx (path or bare flag)",
    )
    run_p.add_argument(
        "--scorer", choices=["stub", "server", "keyword", "vouchpilot"], default="stub"
    )
    run_p.add_argument("--offset", type=int, default=0, help="Start at this canonical row index")
    run_p.add_argument("--limit", type=int, default=0, help="Max rows (0 = all)")
    run_p.add_argument("--endpoint", default="http://localhost:8080")
    run_p.add_argument(
        "--exemplars", default=None, help="Exemplar index prefix (loads PREFIX_exemplars.json)"
    )
    run_p.add_argument("--k", type=int, default=5)
    run_p.add_argument("--margin", type=float, default=0.15)
    run_p.add_argument(
        "--max-challenge-rate",
        type=float,
        default=1.0,
        help="Max share of rows the challenger may re-check (lowest-margin first)",
    )
    run_p.add_argument(
        "--workers",
        type=int,
        default=1,
        help="Parallel row workers (match llama-server slots; >1 is faster but "
        "server batching may flip close calls: use 1 for bit-identical repro)",
    )
    run_p.add_argument(
        "--escalation-endpoint", default=None, help="LLM server URL for challenger escalation"
    )
    run_p.add_argument("--calibrator", default=None, help="Calibrator JSON to rescale confidence")
    run_p.set_defaults(func=cmd_run)

    eval_p = sub.add_parser("evaluate", help="Score predictions against gold labels")
    eval_p.add_argument("--gold", required=True)
    eval_p.add_argument("--pred", required=True)
    eval_p.add_argument("--report", required=True)
    eval_p.add_argument(
        "--calibrator", default=None, help="Calibrator JSON for ECE/coverage/pairwise report"
    )
    eval_p.set_defaults(func=cmd_evaluate)

    cal_p = sub.add_parser("calibrate", help="Fit temperature on predictions")
    cal_p.add_argument("--pred", required=True)
    cal_p.add_argument("--gold", required=True)
    cal_p.add_argument("--out", required=True)
    cal_p.set_defaults(func=cmd_calibrate)

    gold_p = sub.add_parser("gold", help="Generate a synthetic gold dataset")
    gold_p.add_argument("--n", type=int, default=200)
    gold_p.add_argument("--seed", type=int, default=7)
    gold_p.add_argument("--out", required=True)
    gold_p.set_defaults(func=cmd_gold)

    doc_p = sub.add_parser("doctor", help="Environment and dependency checks")
    doc_p.add_argument("--endpoint", default="http://localhost:8080")
    doc_p.add_argument("--weights", default=None)
    doc_p.set_defaults(func=cmd_doctor)

    agent_p = sub.add_parser("agent", help="Agentic loop: inspect, classify, challenge, approve, export")
    agent_sub = agent_p.add_subparsers(dest="agent_command", required=True)
    run_a = agent_sub.add_parser("run", help="Full loop with human approval gate")
    run_a.add_argument("--input", required=True)
    run_a.add_argument("--out", required=True)
    run_a.add_argument("--xlsx", nargs="?", const="auto", default=None)
    run_a.add_argument(
        "--scorer", choices=["stub", "server", "keyword", "vouchpilot"], default="stub"
    )
    run_a.add_argument("--offset", type=int, default=0)
    run_a.add_argument("--limit", type=int, default=0)
    run_a.add_argument("--endpoint", default="http://localhost:8080")
    run_a.add_argument("--exemplars", default=None)
    run_a.add_argument("--k", type=int, default=5)
    run_a.add_argument("--margin", type=float, default=0.15)
    run_a.add_argument("--max-challenge-rate", type=float, default=1.0)
    run_a.add_argument("--workers", type=int, default=1)
    run_a.add_argument("--escalation-endpoint", default=None)
    run_a.add_argument("--calibrator", default=None)
    run_a.add_argument(
        "--gate", default="auto",
        help="Human approval gate: auto | cli | file:PATH-to-decisions.json",
    )
    run_a.add_argument("--pending-out", default=None, help="Also write the review queue JSON")
    run_a.set_defaults(func=cmd_agent_run)

    tmpl_a = agent_sub.add_parser(
        "review-template", help="Emit pending.json review queue without exporting"
    )
    tmpl_a.add_argument("--input", required=True)
    tmpl_a.add_argument("--out", required=True)
    tmpl_a.add_argument(
        "--scorer", choices=["stub", "server", "keyword", "vouchpilot"], default="stub"
    )
    tmpl_a.add_argument("--endpoint", default="http://localhost:8080")
    tmpl_a.add_argument("--exemplars", default=None)
    tmpl_a.add_argument("--k", type=int, default=5)
    tmpl_a.add_argument("--margin", type=float, default=0.15)
    tmpl_a.add_argument("--max-challenge-rate", type=float, default=1.0)
    tmpl_a.add_argument("--workers", type=int, default=1)
    tmpl_a.add_argument("--escalation-endpoint", default=None)
    tmpl_a.set_defaults(func=cmd_agent_template)

    audit_p = sub.add_parser("audit", help="Retro-audit: recorded vs predicted voucher types")
    audit_p.add_argument("--input", required=True, help="xlsx WITH a voucher-type column")
    audit_p.add_argument("--voucher-col", default="voucher_type")
    audit_p.add_argument("--pred", required=True)
    audit_p.add_argument("--report", required=True)
    audit_p.set_defaults(func=cmd_audit)

    robust_p = sub.add_parser("robust", help="Perturbation sweeps for O4/O5 robustness")
    robust_p.add_argument("--gold", required=True, help="Gold prefix (PREFIX.xlsx + PREFIX_labels.json)")
    robust_p.add_argument("--rates", default="0,0.1,0.2,0.3")
    robust_p.add_argument("--rename-rates", default="0,0.25,0.5")
    robust_p.add_argument("--scorer", choices=["stub", "keyword"], default="keyword")
    robust_p.add_argument("--seed", type=int, default=7)
    robust_p.add_argument("--out", required=True)
    robust_p.set_defaults(func=cmd_robust)

    intake_p = sub.add_parser("intake", help="Convert PDF/image/CSV/XLSX documents to rows XLSX")
    intake_p.add_argument("--input", required=True)
    intake_p.add_argument("--out", required=True)
    intake_p.add_argument("--langs", default="hin,eng",
                          help="OCR languages, comma separated (needs tesseract binary)")
    intake_p.set_defaults(func=cmd_intake)
    return parser


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    start = time.time()
    code = args.func(args)
    _ = start
    return code


if __name__ == "__main__":
    sys.exit(main())
